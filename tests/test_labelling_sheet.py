"""Generator tests — the Definitions tab must carry LABEL_SET_v9.md verbatim."""

from pathlib import Path

from openpyxl import load_workbook

from router.build_labelling_sheet import (
    DEFINITIONS_MD,
    header,
    regenerate_definitions,
    write_sheet,
)
from router.labels import load_labels


def test_definitions_cell_is_md_verbatim(tmp_path: Path):
    out = tmp_path / "sheet.xlsx"
    write_sheet(out, load_labels(), [["T-001", "T-001.pdf", "excerpt"]])
    wb = load_workbook(out)
    cell = wb["Definitions"]["A2"].value
    assert cell == DEFINITIONS_MD.read_text(encoding="utf-8")


def test_label_tab_header_order():
    labels = load_labels()
    hdr = header(labels)
    assert hdr == ["id", "file name", "excerpt", *labels, "unsure", "notes"]
    assert len(hdr) == 27


def test_label_cells_are_literal_false(tmp_path: Path):
    labels = load_labels()
    out = tmp_path / "sheet.xlsx"
    write_sheet(out, labels, [["T-001", "T-001.pdf", "excerpt"]])
    wb = load_workbook(out)
    ws = wb["Label"]
    row = [c.value for c in ws[2]]
    assert row[0] == "T-001"
    assert all(v == "False" for v in row[3:25])
    assert row[25] is None and row[26] is None


def test_regen_definitions_leaves_label_tab_untouched(tmp_path: Path):
    labels = load_labels()
    out = tmp_path / "sheet.xlsx"
    write_sheet(out, labels, [["T-001", "T-001.pdf", "excerpt"]])
    wb = load_workbook(out)
    ws = wb["Label"]
    ws.cell(row=2, column=4, value="TRUE")  # an owner tick
    wb.save(out)

    regenerate_definitions(out)
    wb2 = load_workbook(out)
    assert wb2["Definitions"]["A2"].value == DEFINITIONS_MD.read_text(encoding="utf-8")
    assert wb2["Label"].cell(row=2, column=4).value == "TRUE"  # untouched
